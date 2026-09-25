#!/usr/bin/env python3
"""
Replace the image on Ch2 slide 4 ("手把手寫出 5 大壓測型態：stages 腳本範例") with a
native slide: text boxes, code boxes and native XY line charts plotted from the same
`stages` examples the codelab uses (see generate_traffic_pattern_figures.py).

Starts from the deck as committed in 601893d (the image-only version) so it can be
re-run; everything else in the deck is left untouched.
"""
import io
import os
import subprocess
import sys

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from add_metric_reading_slides import FONT, MONO, S, add_box, add_card_text, add_text, px, rgb  # noqa: E402
from generate_traffic_pattern_figures import SMOKE_SECONDS, SMOKE_VUS, STAGES, points  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "k6", "slides", "Ch2_Scientific_k6_Traffic_Modeling.pptx")
BASELINE_COMMIT = "601893d"
SLIDE_INDEX = 3  # slide 4

BG = "071A21"
CARD = "0C2530"
CARD_BORDER = "1E4450"
WHITE = "FFFFFF"
TEXT = "D7E3EA"
MUTED = "8FA6B2"
CYAN = "38BDF8"
MAGENTA = "E879F9"
AMBER = "FBBF24"
CODE_BG = "04121A"
KEY = "93C5FD"
NUM = "FBBF24"
STR = "86EFAC"


def fmt(seconds):
    if seconds % 3600 == 0:
        return f"{seconds // 3600}h"
    if seconds % 60 == 0:
        return f"{seconds // 60}m"
    return f"{seconds}s"


def stage_runs(stages, per_line):
    """Code lines for a stages array, `per_line` stages per line, coloured runs."""
    lines = [[("stages: [", TEXT)]]
    for i in range(0, len(stages), per_line):
        row = [("  ", TEXT)]
        for dur, tgt in stages[i:i + per_line]:
            row += [("{ duration: ", TEXT), (f"'{fmt(dur)}'", STR), (", target: ", TEXT), (str(tgt), NUM), (" }, ", TEXT)]
        lines.append(row)
    lines.append([("]", TEXT)])
    return lines


def mini_chart(slide, x, y, w, h, xs, ys, color):
    cd = XyChartData()
    ser = cd.add_series("VUs")
    for a, b in zip(xs, ys):
        ser.add_data_point(a, b)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, *px(x, y, w, h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    for ax in (ch.value_axis, ch.category_axis):
        ax.has_major_gridlines = False
        ax.major_tick_mark = XL_TICK_MARK.NONE
        ax.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        ax.format.line.color.rgb = rgb("2A4A56")
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = max(ys) * 1.15
    ch.category_axis.minimum_scale = 0
    ch.category_axis.maximum_scale = xs[-1]
    s = ch.plots[0].series[0]
    s.smooth = False
    s.format.line.color.rgb = rgb(color)
    s.format.line.width = Pt(2.5)
    cs = ch._chartSpace
    for parent in (cs, cs.chart.plotArea):
        sppr = parent.find(qn("c:spPr"))
        if sppr is None:
            sppr = parent.makeelement(qn("c:spPr"), {})
            anchor = parent.find(qn("c:txPr")) if parent is cs else None
            if parent is cs and anchor is None:
                anchor = cs.find(qn("c:externalData"))
            if anchor is not None:
                anchor.addprevious(sppr)
            else:
                parent.append(sppr)
        for c in list(sppr):
            sppr.remove(c)
        sppr.append(sppr.makeelement(qn("a:noFill"), {}))
        ln = sppr.makeelement(qn("a:ln"), {})
        ln.append(ln.makeelement(qn("a:noFill"), {}))
        sppr.append(ln)


def build(out=DECK):
    blob = subprocess.check_output(["git", "-C", ROOT, "show", f"{BASELINE_COMMIT}:k6/slides/Ch2_Scientific_k6_Traffic_Modeling.pptx"])
    prs = Presentation(io.BytesIO(blob))
    s = prs.slides[SLIDE_INDEX]
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    assert len(pics) == 1
    s.shapes._spTree.remove(pics[0]._element)
    badge = [sh for sh in s.shapes if sh.name == "Codelabs_Link_Badge"][0]._element

    bg = add_box(s, 0, 0, 1376, 768, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 60, 26, 1256, 52, [[("手把手寫出 5 大壓測型態：stages 腳本範例", S(28, WHITE, True))]], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 60, 78, 1256, 28, [[("k6 會在每一格的 duration 內，把 VU 數線性調整到 target——數字與 Codelab 範例一致", S(14, CYAN))]],
             anchor=MSO_ANCHOR.MIDDLE)

    smoke_t, smoke_v = [0, 0, SMOKE_SECONDS, SMOKE_SECONDS], [0, SMOKE_VUS, SMOKE_VUS, 0]
    rows = [
        ("Smoke Test", "冒煙測試", "1 VU・1 分鐘", (smoke_t, smoke_v), CYAN,
         [[("vus: ", TEXT), ("1", NUM), (", duration: ", TEXT), ("'1m'", STR)]]),
        ("Load Test", "常規負載", "爬坡 → 高原 10m → 降載", points(STAGES["load"]), CYAN, stage_runs(STAGES["load"], 2)),
        ("Stress Test", "極限壓力", "每階爬坡＋維持，到 300 VUs", points(STAGES["stress"]), MAGENTA, stage_runs(STAGES["stress"], 2)),
        ("Spike Test", "突發尖峰", "10 秒內暴增 20 倍到 200 VUs", points(STAGES["spike"]), AMBER, stage_runs(STAGES["spike"], 2)),
        ("Soak Test", "浸泡耐久", "40 VUs 固定長跑 4 小時", points(STAGES["soak"]), CYAN, stage_runs(STAGES["soak"], 2)),
    ]
    heights = [66, 100, 160, 120, 100]  # sized to each code block (1 / 4 / 7 / 5 / 4 lines)
    y, gap = 118, 8
    for (name, zh, desc, (xs, ys), color, code), rh in zip(rows, heights):
        add_box(s, 60, y, 1256, rh, fill=CARD, line=CARD_BORDER, radius=0.08)
        add_text(s, 80, y + 6, 200, rh - 12, [
            [(name, S(16, WHITE, True)), ("  " + zh, S(12.5, MUTED, True))],
            [(desc, S(11.5, TEXT))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
        mini_chart(s, 280, y + 8, 330, rh - 16, xs, ys, color)
        cb = add_box(s, 628, y + 10, 674, rh - 20, fill=CODE_BG, line=CARD_BORDER, radius=0.08)
        paras = [[(t, S(13, c, False, MONO)) for t, c in line] for line in code]
        add_card_text(cb, paras, anchor=MSO_ANCHOR.MIDDLE, margin=(16, 0, 12, 0), line_spacing=1.0)
        y += rh + gap

    # keep the badge on top of the new background
    tree = s.shapes._spTree
    tree.remove(badge)
    tree.append(badge)
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else DECK)
