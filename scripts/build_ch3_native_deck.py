#!/usr/bin/env python3
"""
Rebuild k6/slides/Ch3_k6_Quality_Gates.pptx as an all-native deck: text boxes,
shapes, tables and connectors instead of full-slide PNGs, so future edits are
text edits rather than image regeneration.

Only the cover keeps a picture (the illustration, cropped from the original cover).
Slides 5 and 6 come from add_metric_reading_slides.py.

The source images are read from the deck as committed in 601893d (the last
image-only version), so the script can be re-run at any time:

    python3 scripts/build_ch3_native_deck.py
"""
import copy
import io
import os
import subprocess
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import add_metric_reading_slides as base  # noqa: E402
from add_metric_reading_slides import (  # noqa: E402
    FONT, MONO, S, add_box, add_card_text, add_text, blank_layout, light_header, px, rgb,
    L_ACCENT, L_ACCENT_BG, L_ACCENT_BORDER, L_BG, L_BODY, L_BORDER, L_MUTED, L_TITLE,
    T_AMBER, T_BG, T_CYAN, T_DIM, T_GREEN, T_RED, T_TEXT,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "k6", "slides", "Ch3_k6_Quality_Gates.pptx")
BASELINE_COMMIT = "601893d"  # last image-only version of the deck
CODELAB_URL = base.CODELAB_URL

RED_C = "DC2626"
RED_BG = "FEF2F2"
RED_BORDER = "FECACA"
GREEN_C = "15803D"
GREEN_BG = "F0FDF4"
GREEN_BORDER = "BBF7D0"
BLUE_C = "2563EB"
AMBER_C = "B45309"
SLATE_BG = "F1F5F9"


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def baseline_deck():
    blob = subprocess.check_output(["git", "-C", ROOT, "show", f"{BASELINE_COMMIT}:k6/slides/Ch3_k6_Quality_Gates.pptx"])
    return Presentation(io.BytesIO(blob))


def badge_element(prs):
    for s in prs.slides:
        for shp in s.shapes:
            if shp.name == "Codelabs_Link_Badge":
                return copy.deepcopy(shp._element)
    raise RuntimeError("badge not found")


def add_badge(slide, el):
    el = copy.deepcopy(el)
    rid = slide.part.relate_to(CODELAB_URL, RT.HYPERLINK, is_external=True)
    for h in el.iter(qn("a:hlinkClick")):
        h.set(qn("r:id"), rid)
    slide.shapes._spTree.append(el)


def arrow(slide, x1, y1, x2, y2, color=L_MUTED, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *px(x1, y1, 0, 0)[:2], *px(x2, y2, 0, 0)[:2])
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def line(slide, x1, y1, x2, y2, color=L_MUTED, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *px(x1, y1, 0, 0)[:2], *px(x2, y2, 0, 0)[:2])
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def code_box(slide, x, y, w, h, lines, size=14, numbered=False, dark=True, anchor=MSO_ANCHOR.TOP, pad=(22, 16, 16, 12)):
    """lines: list of list of (text, color) runs."""
    box = add_box(slide, x, y, w, h, fill=T_BG if dark else "FFFFFF", line="334155" if dark else L_BORDER, radius=0.03)
    paras = []
    for i, runs in enumerate(lines, 1):
        p = []
        if numbered:
            p.append((f"{i:>2}  ", T(size, T_DIM if dark else L_MUTED, font=MONO)))
        for t, c in runs:
            # a run is (text, colour) or (text, full style dict)
            p.append((t, c if isinstance(c, dict) else T(size, c, font=MONO)))
        paras.append(p)
    add_card_text(box, paras, anchor=anchor, margin=pad, line_spacing=1.12)
    return box


BODY_SCALE = 1.2  # body text on a 13.3"-equivalent canvas reads small at the original sizes


def T(size, color, bold=False, font=FONT):
    return S(round(size * BODY_SCALE, 1), color, bold, font)


def new_slide(prs, badge):
    s = prs.slides.add_slide(blank_layout(prs))
    add_badge(s, badge)
    return s


# ---------------------------------------------------------------------------
# slides
# ---------------------------------------------------------------------------
def s1_cover(prs, badge, cover_png):
    s = prs.slides.add_slide(blank_layout(prs))
    add_box(s, 0, 0, 1376, 768, fill="FFFFFF", shape=MSO_SHAPE.RECTANGLE)
    # right-hand illustration, cropped from the original cover (no text in this region)
    im = Image.open(io.BytesIO(cover_png)).convert("RGB")
    sx = im.width / 1376
    crop = im.crop((int(740 * sx), 0, im.width, int(525 * sx)))
    buf = io.BytesIO()
    crop.save(buf, "PNG")
    buf.seek(0)
    s.shapes.add_picture(buf, *px(740, 0, 636, 525))
    add_text(s, 70, 110, 660, 250, [
        [("讓數據說話：", S(46, L_TITLE, True))],
        [("從指標解讀到", S(46, L_TITLE, True))],
        [("CI/CD 效能品質防線", S(46, L_TITLE, True))],
    ], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    add_text(s, 72, 380, 660, 90, [
        [("Module 3：效能指標解讀與", S(24, L_MUTED, True))],
        [("SLO 門檻自動化（22～25 分鐘）", S(24, L_MUTED, True))],
    ], line_spacing=1.1)
    cards = [
        ("觀測力 Observe", "讀懂 k6 結尾摘要、RED Method 與延遲拆解", BLUE_C),
        ("攔截規則 Define", "擺脫平均值陷阱，自訂指標與 p95 / p99 門檻", AMBER_C),
        ("自動化機關 Automate", "Thresholds 觸發 Exit Code 99，守住 CI/CD", "6D28D9"),
    ]
    for i, (t, d, c) in enumerate(cards):
        x = 60 + i * 425
        add_box(s, x, 560, 405, 150, fill="FFFFFF", line=L_BORDER, radius=0.08)
        add_text(s, x + 24, 578, 360, 116, [
            [(t, T(19, c, True))],
            [(d, T(14.5, L_BODY))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=6)
    add_badge(s, badge)
    return s


def s2_red(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "一眼看穿系統健康度：k6 摘要與 RED Method",
                 "結尾摘要裡最重要的三行，剛好對應微服務的 RED Method：Rate、Errors、Duration")
    W, D = T_TEXT, T_DIM
    mk = lambda t, c: ("   ◀ " + t, T(13.5, c, True, MONO))  # noqa: E731
    lines = [
        [("█ THRESHOLDS", T(13.5, W, True, MONO))],
        [("  http_req_duration", T(13.5, D, font=MONO))],
        [("  ✓ 'p(95)<500' p(95)=312.4ms", T(13.5, T_GREEN, font=MONO))],
        [("  http_req_failed", T(13.5, D, font=MONO))],
        [("  ✓ 'rate<0.01' rate=0.20%", T(13.5, T_GREEN, font=MONO))],
        [(" ", T(6, D, font=MONO))],
        [("█ TOTAL RESULTS", T(13.5, W, True, MONO))],
        [("  HTTP", T(13.5, T_CYAN, True, MONO))],
        [("  http_req_duration: med=121ms p(95)=312ms", T(13.5, W, font=MONO)), mk("Duration", T_AMBER)],
        [("  http_req_failed..: 0.20%  6 out of 3000", T(13.5, W, font=MONO)), mk("Errors", T_RED)],
        [("  http_reqs........: 3000   49.9/s", T(13.5, W, font=MONO)), mk("Rate", T_CYAN)],
        [("  EXECUTION", T(13.5, T_CYAN, True, MONO))],
        [("  iterations.......: 3000   49.9/s", T(13.5, W, font=MONO))],
        [("  vus..............: 50     min=50 max=50", T(13.5, W, font=MONO))],
        [("  NETWORK", T(13.5, T_CYAN, True, MONO))],
        [("  data_received....: 12 MB  200 kB/s", T(13.5, W, font=MONO))],
    ]
    code_box(s, 60, 132, 690, 560, lines, anchor=MSO_ANCHOR.MIDDLE)
    cards = [
        ("Rate（吞吐量）", "http_reqs", "每秒發出多少請求（RPS），衡量系統承載的流量", BLUE_C),
        ("Errors（錯誤率）", "http_req_failed", "非 2xx / 3xx 回應的比例，直接反映可用性", RED_C),
        ("Duration（延遲）", "http_req_duration", "送出請求到收完回應的耗時；看 p95 / p99，不看平均", AMBER_C),
    ]
    y = 132
    for t, m, d, c in cards:
        add_box(s, 774, y, 542, 152, fill="FFFFFF", line=L_BORDER, radius=0.08)
        add_text(s, 796, y + 10, 500, 132, [
            [(t, T(18, c, True))],
            [(m, T(15, L_TITLE, True, MONO))],
            [(d, T(13.5, L_BODY))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=4)
        y += 152 + 10
    note = add_box(s, 774, y, 542, 692 - y, fill=L_ACCENT_BG, line=L_ACCENT_BORDER, radius=0.12)
    add_card_text(note, [
        [("Saturation（飽和度）", T(14, L_ACCENT, True))],
        [("看 vus 是否碰到 vus_max、有沒有 dropped_iterations；延遲慢在哪一段，見第 6 頁的延遲拆解", T(12.5, L_TITLE))],
    ], anchor=MSO_ANCHOR.MIDDLE, margin=(18, 6, 18, 6), space_after=2)
    return s


def s3_average(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "擺脫平均值謊言：專注於 p95 / p99 尾端延遲",
                 "平均值會把少數極慢的請求「攤平」——而那往往是正在結帳的使用者")
    # waffle: 100 requests
    add_box(s, 60, 132, 470, 560, fill="FFFFFF", line=L_BORDER, radius=0.05)
    add_text(s, 84, 146, 430, 34, [[("100 個請求", T(18, L_TITLE, True))]], anchor=MSO_ANCHOR.MIDDLE)
    cell, gap, gx, gy = 30, 8, 108, 196
    for i in range(100):
        r, c = divmod(i, 10)
        slow = i == 99
        add_box(s, gx + c * (cell + gap), gy + r * (cell + gap), cell, cell,
                fill=RED_C if slow else "94A3B8", radius=0.2)
    add_text(s, 84, 584, 430, 96, [
        [("■ ", T(14, "94A3B8", True)), ("99 個請求：10ms", T(14.5, L_BODY))],
        [("■ ", T(14, RED_C, True)), ("1 個請求：10,000ms（資料庫鎖卡住的結帳）", T(14.5, L_BODY))],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    # stat tiles
    tiles = [
        ("平均值", "109.9 ms", "(99 × 10 + 10,000) ÷ 100——看起來很快", L_ACCENT, L_ACCENT_BG, L_ACCENT_BORDER),
        ("最慢的那 1 個", "10,000 ms", "平均值完全看不出來，只有看尾端才抓得到", RED_C, RED_BG, RED_BORDER),
    ]
    x = 554
    for label, big, sub, c, bg, bd in tiles:
        add_box(s, x, 132, 371, 196, fill=bg, line=bd, radius=0.08)
        add_text(s, x + 22, 146, 327, 168, [
            [(label, T(15, c, True))],
            [(big, T(40, L_TITLE, True))],
            [(sub, T(12.5, L_BODY))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
        x += 371 + 20
    add_box(s, 554, 344, 762, 348, fill="FFFFFF", line=L_BORDER, radius=0.05)
    bullets = [
        ("平均值會掩蓋極端值", "少數極慢的請求被多數快請求攤平，報表上看不到真實痛點。"),
        ("p(95)：常用的發版門檻", "95% 的請求都要在門檻內，允許少量突波但不失控。"),
        ("p(99)：尾端延遲的防線", "捕捉 GC 停頓、DB Lock、連線排隊這類只影響少數請求的問題。"),
        ("微服務會放大尾端延遲", "一個頁面呼叫 10 個服務、每個有 1% 變慢，就有 1 − 0.99¹⁰ ≈ 9.6% 的頁面會卡。"),
    ]
    add_text(s, 578, 364, 716, 312, [p for t, d in bullets for p in (
        [("● ", T(14, L_ACCENT, True)), (t, T(16, L_TITLE, True))],
        [("   " + d, T(14, L_BODY))],
    )], anchor=MSO_ANCHOR.MIDDLE, space_after=6)
    return s


def s4_builtin(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "k6 內建核心指標全覽 (Built-in Metrics)",
                 "不用寫任何程式，每次測試結束都會自動輸出，也是設定 Thresholds 的主要標的")
    cats = [
        ("HTTP 狀態", "HTTP Category", BLUE_C, [
            ("http_reqs", "總請求數與 RPS（RED：Rate）"),
            ("http_req_failed", "失敗率（RED：Errors）"),
            ("http_req_duration", "請求耗時（RED：Duration）"),
        ]),
        ("流量與網路", "Network Category", "0F766E", [
            ("data_sent", "送出的資料量"),
            ("data_received", "收到的資料量：回應是否太肥、網卡是否快滿"),
        ]),
        ("執行進度", "Execution Category", "6D28D9", [
            ("vus", "當下的虛擬用戶數"),
            ("vus_max", "虛擬用戶數上限"),
            ("iterations", "完整跑完腳本的次數"),
        ]),
    ]
    for i, (name, en, c, items) in enumerate(cats):
        x = 60 + i * 426
        add_box(s, x, 132, 404, 470, fill="FFFFFF", line=L_BORDER, radius=0.05)
        add_text(s, x + 24, 148, 356, 70, [
            [(name, T(22, c, True))],
            [(en, T(13, L_MUTED))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
        line(s, x + 24, 230, x + 380, 230, color=L_BORDER, width=1)
        yy = 250
        for m, d in items:
            add_text(s, x + 24, yy, 356, 100, [
                [(m, T(17, L_TITLE, True, MONO))],
                [(d, T(13.5, L_BODY))],
            ], space_after=3)
            yy += 112
    bar = add_box(s, 60, 622, 1256, 76, fill=L_ACCENT_BG, line=L_ACCENT_BORDER, radius=0.2)
    add_card_text(bar, [[("HTTP 三兄弟直接對應 RED；執行進度這一類回答的是「這次壓測本身有沒有照計畫執行」", T(15, L_TITLE, True))]],
                  anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    return s


def s7_custom(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "監控業務專屬數據：4 大自訂指標型態",
                 "內建指標只看得到 HTTP；業務語意（結帳次數、佇列長度、業務失敗率）要自己宣告")
    kinds = [
        ("Counter", "累積計數器（像里程表）", "只會往上加總的數字", "累積業務發生次數，例如成功結帳數", "checkout_count.add(1);", BLUE_C),
        ("Gauge", "即時數值（像時速表）", "只保留最新的一個值", "追蹤佇列長度、連線數等即時狀態", "queue_size.add(currentLen);", "0F766E"),
        ("Rate", "成功 / 失敗比率（像紅綠燈）", "統計 true（非零）出現的比例", "業務邏輯失敗率，例如付款失敗比例", "biz_fail_rate.add(!isSuccess);", RED_C),
        ("Trend", "耗時分佈（像趨勢圖）", "可計算 avg / med / p95 / p99", "統計特定區塊的耗時分佈，例如 DB 查詢", "db_query_time.add(duration);", AMBER_C),
    ]
    for i, (name, alias, what, use, code, c) in enumerate(kinds):
        x = 60 + (i % 2) * 638
        y = 132 + (i // 2) * 290
        add_box(s, x, y, 618, 272, fill="FFFFFF", line=L_BORDER, radius=0.05)
        add_text(s, x + 24, y + 14, 570, 150, [
            [(name, T(24, c, True, MONO)), ("   " + alias, T(15, L_MUTED, True))],
            [(what, T(15, L_TITLE, True))],
            [("適用：" + use, T(14, L_BODY))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=6)
        cb = add_box(s, x + 24, y + 180, 570, 70, fill=T_BG, line="334155", radius=0.12)
        add_card_text(cb, [[(code, T(16, T_TEXT, font=MONO))]], anchor=MSO_ANCHOR.MIDDLE, margin=(20, 0, 20, 0))
    return s


def s8_assertions(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "釐清斷言選擇：k6 斷言三部曲完整對照",
                 "check() 看單次回應、thresholds 看整場統計、expect() 是 BDD 風格語法")
    rows = [
        ("", "check()\n請求級軟斷言", "thresholds\n全域效能門檻", "expect()\nBDD 風格斷言"),
        ("作用層級", "單次請求回應\n（Status / Body）", "整場測試的聚合統計\n（例如 p95）", "程式碼單元 / 局部邏輯"),
        ("失敗後果", "記錄失敗比率，測試繼續\n不影響 Exit Code", "門檻未達標\n→ Exit Code 99", "拋出例外\n中斷當前迭代"),
        ("核心用途", "驗證回應內容正確", "CI/CD 自動卡關\n（Performance Gate）", "類 Jest / Chai 的\n流暢斷言語法"),
    ]
    head_fill = [SLATE_BG, "FEF9C3", RED_BG, "F3E8FF"]
    body_fill = ["F8FAFC", "FEFCE8", "FFF5F5", "FAF5FF"]
    tbl = s.shapes.add_table(4, 4, *px(60, 132, 1256, 440)).table
    widths = [220, 345, 345, 346]
    for j, w in enumerate(widths):
        tbl.columns[j].width = px(0, 0, w, 0)[2]
    for i in range(4):
        tbl.rows[i].height = px(0, 0, 0, 110)[3]
        for j in range(4):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill[j] if i == 0 else body_fill[j])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.clear()
            txt = rows[i][j]
            parts = txt.split("\n") if txt else [""]
            for k, part in enumerate(parts):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = part
                head = i == 0 and k == 0
                r.font.size = Pt(22 if head else (18 if (i == 0 or j == 0) else 17.5))
                r.font.bold = i == 0 or j == 0
                r.font.color.rgb = rgb(L_TITLE if (i == 0 or j == 0) else L_BODY)
                r.font.name = MONO if (head and j > 0) else FONT
                rpr = r._r.get_or_add_rPr()
                ea = rpr.makeelement(qn("a:ea"), {"typeface": MONO if (head and j > 0) else FONT})
                rpr.append(ea)
    bar = add_box(s, 60, 594, 1256, 104, fill=T_BG, line="334155", radius=0.12)
    add_card_text(bar, [
        [("最佳實踐：", T(16, T_AMBER, True)), ("用 check() 驗證回應正確性，再在 thresholds 宣告", T(16, T_TEXT))],
        [("'checks': ['rate>0.99']", T(17, T_CYAN, True, MONO)), ("   ——讓資料錯誤也能讓 CI/CD 卡關", T(16, T_TEXT))],
    ], anchor=MSO_ANCHOR.MIDDLE, margin=(28, 0, 28, 0), space_after=4)
    return s


def s9_tags(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "精準控制各 API 門檻：Thresholds 與 Tags 過濾",
                 "替請求貼上標籤，就能對不同端點設定不同的門檻，次要端點變慢不會誤殺整條發版流水線")
    raw = add_box(s, 60, 134, 610, 70, fill=SLATE_BG, line="CBD5E1", radius=0.2)
    add_card_text(raw, [[("原始 http_req_duration 數據（所有請求）", T(16, L_TITLE, True))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    for x, tag, gate, c in ((60, "{api_type:critical}", "p(95) < 300ms", RED_C), (410, "{api_type:background}", "p(95) < 5000ms", "0F766E")):
        arrow(s, x + 130, 206, x + 130, 262, color=c)
        tb = add_box(s, x, 264, 260, 86, fill="FFFFFF", line=c, radius=0.12, line_w=1.75)
        add_card_text(tb, [[("Tag", T(13, L_MUTED, True))], [(tag, T(15, c, True, MONO))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        arrow(s, x + 130, 352, x + 130, 470, color=c)
        gb = add_box(s, x, 472, 260, 92, fill=c, line=c, radius=0.12)
        add_card_text(gb, [[("門檻", T(13, "E0E7FF", True))], [(gate, T(17, "FFFFFF", True, MONO))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    grp = add_box(s, 60, 590, 610, 104, fill="FFF7ED", line="FED7AA", radius=0.12)
    add_card_text(grp, [
        [("Group 也能過濾：", T(15, AMBER_C, True)), ("{group:::01_核心結帳交易}", T(15, L_TITLE, True, MONO))],
        [("用 group() 包起來的整段流程，可以一起設門檻", T(14, L_BODY))],
    ], anchor=MSO_ANCHOR.MIDDLE, margin=(22, 0, 22, 0), space_after=4)
    G, Y, P, W = T_GREEN, T_AMBER, "C4B5FD", T_TEXT
    lines = [
        [("export const ", P), ("options", W), (" = {", W)],
        [("  thresholds: {", W)],
        [("    'http_req_duration", G), ("{api_type:critical}", "FCA5A5"), ("':", G)],
        [("      ['p(95)<300'],", Y)],
        [("", W)],
        [("    'http_req_duration", G), ("{api_type:background}", "5EEAD4"), ("':", G)],
        [("      ['p(95)<5000'],", Y)],
        [("", W)],
        [("    'http_req_duration", G), ("{group:::01_核心結帳交易}", "FDBA74"), ("':", G)],
        [("      ['p(95)<1000'],", Y)],
        [("  },", W)],
        [("};", W)],
        [("", W)],
        [("// 發送請求時貼上標籤", T_DIM)],
        [("http.get(url, { tags: { api_type: 'critical' } });", W)],
    ]
    code_box(s, 694, 134, 622, 560, lines, size=14.5, anchor=MSO_ANCHOR.MIDDLE)
    return s


def s10_exit_code(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "CI/CD 自動卡關實務：Exit Code 99 傳遞鏈",
                 "k6 門檻失敗會回傳 99，但只要 Exit Code 在途中被吃掉，CI 就會亮綠燈放行")
    panels = [
        (60, "×  錯誤示範：退出碼被遮蔽", RED_C, RED_BG, RED_BORDER,
         ["k6 結束\n（Exit Code 99）", "後面接著 echo 等指令\n最後一個指令成功 → 回傳 0", "CI/CD Pipeline\n顯示 PASS"],
         "k6 run test.js; echo \"done\"", "未爆彈混入生產環境"),
        (698, "✓  正確示範：明確捕捉與傳遞", GREEN_C, GREEN_BG, GREEN_BORDER,
         ["k6 結束\n（Exit Code 99）", "code=$?\n立即捕捉退出碼", "exit $code\n把 99 傳下去", "CI/CD Pipeline\n顯示 BLOCKED"],
         "k6 run test.js; code=$?; echo \"done\"; exit $code", "成功阻擋有效能缺陷的 PR 合併"),
    ]
    for x, head, c, bg, bd, steps, cmd, result in panels:
        add_box(s, x, 132, 618, 566, fill=bg, line=bd, radius=0.04)
        add_text(s, x + 24, 144, 570, 40, [[(head, T(19, c, True))]], anchor=MSO_ANCHOR.MIDDLE)
        n = len(steps)
        bh, gap = (72, 22) if n == 4 else (86, 34)
        y = 196
        for i, st in enumerate(steps):
            b = add_box(s, x + 150, y, 318, bh, fill="FFFFFF", line=c, radius=0.12, line_w=1.5)
            add_card_text(b, [[(t, T(14.5 if k == 0 else 13, L_TITLE if k == 0 else L_BODY, k == 0))] for k, t in enumerate(st.split("\n"))],
                          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            if i < n - 1:
                arrow(s, x + 309, y + bh + 2, x + 309, y + bh + gap - 2, color=c)
            y += bh + gap
        cb = add_box(s, x + 24, 562, 570, 50, fill=T_BG, line="334155", radius=0.15)
        add_card_text(cb, [[("$ " + cmd, T(13, T_TEXT, font=MONO))]], anchor=MSO_ANCHOR.MIDDLE, margin=(16, 0, 16, 0))
        add_text(s, x + 24, 626, 570, 50, [[(result, T(17, c, True))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    return s


def s11_anatomy(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "模組綜合實戰：門檻、熔斷與自訂指標",
                 "把自訂指標、Tags、Thresholds 與 abortOnFail 熔斷組裝成一台自動化效能測試引擎")
    K, G, Y, W = "C4B5FD", T_GREEN, T_AMBER, T_TEXT
    code = [
        [("import", K), (" http ", W), ("from", K), (" 'k6/http';", G)],
        [("import", K), (" { check } ", W), ("from", K), (" 'k6';", G)],
        [("import", K), (" { Rate } ", W), ("from", K), (" 'k6/metrics';", G)],
        [("", W)],
        [("const", K), (" bizSuccess = ", W), ("new", K), (" Rate(", W), ("'business_transaction_success'", G), (");", W)],
        [("", W)],
        [("export const", K), (" options = {", W)],
        [("  thresholds: {", W)],
        [("    'http_req_duration{api_type:critical}'", G), (": [", W), ("'p(95)<300'", Y), ("],", W)],
        [("    'business_transaction_success'", G), (": [{", W)],
        [("      threshold: ", W), ("'rate>=0.95'", Y), (",", W)],
        [("      abortOnFail: ", W), ("true", K), (",", W)],
        [("      delayAbortEval: ", W), ("'5s'", Y), (",", W)],
        [("    }],", W)],
        [("  },", W)],
        [("};", W)],
        [("", W)],
        [("export default function", K), (" () {", W)],
        [("  const res = http.get(url, { tags: { api_type: ", W), ("'critical'", G), (" } });", W)],
        [("  bizSuccess.add(check(res, { ", W), ("'status 200'", G), (": (r) => r.status === 200 }));", W)],
        [("}", W)],
    ]
    size = 14
    lh = 1.435 * size * BODY_SCALE  # measured line pitch in px for Noto Sans Mono CJK at 1.12 spacing
    top = 150 + 0.8 * lh             # centre of line 1 (box top 132 + 18px padding)
    code_box(s, 60, 132, 900, 566, code, size=size, numbered=True, pad=(22, 18, 16, 12))
    callouts = [
        (3, 5, "宣告 Rate 自訂指標", "Line 3・5"),
        (9, 9, "Tags 過濾門檻\n未達標 → Exit Code 99", "Line 9"),
        (10, 14, "abortOnFail 即時熔斷\ndelayAbortEval 暖機寬限", "Line 10–14"),
        (19, 20, "貼標籤＋check 結果\n寫入業務成功率", "Line 19・20"),
    ]
    ys = [140, 282, 424, 566]
    for (a, b, text, tag), y in zip(callouts, ys):
        cb = add_box(s, 1000, y, 316, 118, fill="FFFFFF", line=L_ACCENT_BORDER, radius=0.12, line_w=1.5)
        add_card_text(cb, [[(tag, T(12, L_MUTED, True, MONO))]] + [[(t, T(15, L_TITLE, True))] for t in text.split("\n")],
                      anchor=MSO_ANCHOR.MIDDLE, margin=(18, 0, 12, 0), space_after=1)
        target_y = top + ((a + b) / 2 - 1) * lh
        line(s, 962, target_y, 998, y + 59, color=L_ACCENT, width=1.5)
    return s


def s12_lab(prs, badge):
    s = new_slide(prs, badge)
    light_header(s, "Module 3 實作練習：打造你的自動化防線 (Lab Guide)",
                 "同一支腳本 k6/demos/ch3_quality_gates_exit99.js，靠環境變數切出三種情境（目標：QuickPizza）")
    tasks = [
        ("任務一：通過情境", "k6 run -e FAIL_SLO=false k6/demos/ch3_quality_gates_exit99.js ; echo \"CI Exit Code: $?\"",
         "觀察 Tags / Group 各自的 p95 與四大自訂指標；用 5 步驟判讀這份摘要 → Exit Code 0"),
        ("任務二：門檻違規", "k6 run -e FAIL_SLO=true k6/demos/ch3_quality_gates_exit99.js ; echo \"CI Exit Code: $?\"",
         "只有 critical 端點打紅叉，background 仍是綠的 → Exit Code 99"),
        ("任務三：即時熔斷", "k6 run -e ABORT_TEST=true k6/demos/ch3_quality_gates_exit99.js ; echo \"CI Exit Code: $?\"",
         "abortOnFail 觸發，測試在幾秒內提前中止 → Exit Code 99"),
    ]
    y = 134
    for t, cmd, obs in tasks:
        add_box(s, 60, y, 1256, 156, fill="FFFFFF", line=L_BORDER, radius=0.06)
        add_text(s, 84, y + 12, 1208, 34, [[(t, T(18, L_ACCENT, True))]], anchor=MSO_ANCHOR.MIDDLE)
        cb = add_box(s, 84, y + 52, 1208, 46, fill=T_BG, line="334155", radius=0.15)
        add_card_text(cb, [[("$ " + cmd, T(13.5, T_TEXT, font=MONO))]], anchor=MSO_ANCHOR.MIDDLE, margin=(16, 0, 16, 0))
        add_text(s, 84, y + 106, 1208, 40, [[("觀察：", T(14, AMBER_C, True)), (obs, T(14, L_BODY))]], anchor=MSO_ANCHOR.MIDDLE)
        y += 156 + 12
    banner = add_box(s, 360, 648, 656, 58, fill=T_BG, line=T_AMBER, radius=0.3, line_w=2)
    add_card_text(banner, [[("準備好你的 Terminal，我們實作見！", T(20, T_AMBER, True))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    return s


# ---------------------------------------------------------------------------
def build(out=DECK):
    src = baseline_deck()
    badge = badge_element(src)
    cover_png = None
    for shp in src.slides[0].shapes:
        if shp.shape_type == 13:
            cover_png = shp.image.blob
    # start from the baseline deck (keeps masters, theme, slide size) and drop every slide
    prs = baseline_deck()
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)

    s1_cover(prs, badge, cover_png)
    s2_red(prs, badge)
    s3_average(prs, badge)
    s4_builtin(prs, badge)
    base.ch3_summary_sop(prs)       # slide 5 (uses the badge already on slide 1)
    base.ch3_latency_breakdown(prs)  # slide 6
    s7_custom(prs, badge)
    s8_assertions(prs, badge)
    s9_tags(prs, badge)
    s10_exit_code(prs, badge)
    s11_anatomy(prs, badge)
    s12_lab(prs, badge)
    # backgrounds are drawn after the badge on most slides; bring every badge to the front
    for slide in prs.slides:
        tree = slide.shapes._spTree
        for shp in [x for x in slide.shapes if x.name == "Codelabs_Link_Badge"]:
            tree.remove(shp._element)
            tree.append(shp._element)
    prs.save(out)
    print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else DECK)
