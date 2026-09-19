#!/usr/bin/env python3
"""
Rebuild k6/slides/Ch4_Precision_k6_Hybrid_Testing.pptx with native PowerPoint text boxes
for Slide 4 (bottom banner), Slide 5 (3 golden rules cards), and Slide 8 (top-right problem box).
"""
import os
import subprocess
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SW = Emu(16256000)
SH = Emu(9144000)
EMU_PER_PX_X = 16256000 / 1376
EMU_PER_PX_Y = 9144000 / 768

def px(x, y, w, h):
    return (
        int(x * EMU_PER_PX_X),
        int(y * EMU_PER_PX_Y),
        int(w * EMU_PER_PX_X),
        int(h * EMU_PER_PX_Y)
    )

FONT_NAME = "Noto Sans CJK TC"
CODE_FONT = "DejaVu Sans Mono"

def style_box(shape, fill_rgb, line_rgb=None, line_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()

def clean_slide_backgrounds():
    os.makedirs('/tmp/ch4_cleaned', exist_ok=True)
    
    # 1. Slide 4: Inpaint only bottom banner
    s4_img = Image.open('k6/slides/assets/Ch4/slide_4.png').copy()
    draw4 = ImageDraw.Draw(s4_img)
    draw4.rectangle([68, 655, 1308, 740], fill=(53, 36, 36))
    s4_clean = '/tmp/ch4_cleaned/slide_4_clean.png'
    s4_img.save(s4_clean)
    
    # 2. Slide 5: Inpaint 3 sub-boxes in each of the 3 cards
    s5_img = Image.open('k6/slides/assets/Ch4/slide_5.png').copy()
    draw5 = ImageDraw.Draw(s5_img)
    # Card 1
    draw5.rectangle([276, 224, 520, 294], fill=(38, 41, 47))
    draw5.rectangle([276, 300, 520, 440], fill=(30, 32, 38))
    draw5.rectangle([276, 446, 520, 588], fill=(30, 32, 38))
    # Card 2
    draw5.rectangle([566, 224, 809, 294], fill=(38, 41, 47))
    draw5.rectangle([566, 300, 809, 440], fill=(30, 32, 38))
    draw5.rectangle([566, 446, 809, 588], fill=(20, 22, 26))
    # Card 3
    draw5.rectangle([855, 224, 1098, 294], fill=(38, 41, 47))
    draw5.rectangle([855, 300, 1098, 440], fill=(30, 32, 38))
    draw5.rectangle([855, 446, 1098, 588], fill=(30, 32, 38))
    s5_clean = '/tmp/ch4_cleaned/slide_5_clean.png'
    s5_img.save(s5_clean)
    
    # 3. Slide 8: Inpaint top-right window body
    s8_img = Image.open('k6/slides/assets/Ch4/slide_8.png').copy()
    draw8 = ImageDraw.Draw(s8_img)
    draw8.rectangle([950, 67, 1332, 169], fill=(34, 38, 41))
    s8_clean = '/tmp/ch4_cleaned/slide_8_clean.png'
    s8_img.save(s8_clean)
    
    return s4_clean, s5_clean, s8_clean

def build_presentation(s4_clean, s5_clean, s8_clean, output_path):
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]
    
    for i in range(1, 12):
        s = prs.slides.add_slide(blank)
        
        if i == 4:
            s.shapes.add_picture(s4_clean, 0, 0, width=SW, height=SH)
            
            # Bottom banner background
            l, t, w, h = px(68, 655, 1240, 85)
            box4 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
            style_box(box4, RGBColor(53, 36, 36), RGBColor(95, 45, 45), 1)

            # Left red accent stripe
            l_bar, t_bar, w_bar, h_bar = px(68, 655, 5, 85)
            bar4 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l_bar, t_bar, w_bar, h_bar)
            style_box(bar4, RGBColor(237, 79, 70), None)

            # Banner text
            tf4 = box4.text_frame
            tf4.word_wrap = True
            tf4.margin_left = Pt(22)
            tf4.margin_right = Pt(14)
            tf4.margin_top = Pt(6)
            tf4.margin_bottom = Pt(6)
            tf4.vertical_anchor = MSO_ANCHOR.MIDDLE

            p1 = tf4.paragraphs[0]
            p1.alignment = PP_ALIGN.LEFT
            p1.line_spacing = 1.25
            r_icon = p1.add_run()
            r_icon.text = "⚠ 致命陷阱："
            r_icon.font.name = FONT_NAME
            r_icon.font.size = Pt(15)
            r_icon.font.bold = True
            r_icon.font.color.rgb = RGBColor(255, 95, 86)

            r_t1 = p1.add_run()
            r_t1.text = " 產出的腳本為「寫死的硬編碼 (Hardcoded)」！現狀無法直接用於高併發壓測。所有的 Token、Session ID"
            r_t1.font.name = FONT_NAME
            r_t1.font.size = Pt(14.5)
            r_t1.font.color.rgb = RGBColor(241, 245, 249)

            p2 = tf4.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.line_spacing = 1.25
            r_t2 = p2.add_run()
            r_t2.text = "都是錄製當下的死資料。幾千個 VU 用同一個過期 Token 打 API，只會引發大規模 401 Unauthorized 錯誤，結果嚴重失真。"
            r_t2.font.name = FONT_NAME
            r_t2.font.size = Pt(14.5)
            r_t2.font.color.rgb = RGBColor(203, 213, 225)
            
        elif i == 5:
            s.shapes.add_picture(s5_clean, 0, 0, width=SW, height=SH)
            
            # ---------------- CARD 1 (Clean Static Assets) ----------------
            l, t, w, h = px(277, 225, 242, 68)
            h1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(h1, RGBColor(38, 41, 47), RGBColor(70, 75, 86))
            tf = h1.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = "① 剔除靜態與第三方資源\n"
            r.font.name = FONT_NAME
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(241, 245, 249)
            r_sub = p.add_run()
            r_sub.text = "(Clean Static Assets)"
            r_sub.font.name = FONT_NAME
            r_sub.font.size = Pt(11)
            r_sub.font.color.rgb = RGBColor(148, 163, 184)

            # Action Box 1
            l, t, w, h = px(277, 302, 242, 136)
            m1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(m1, RGBColor(30, 32, 38), RGBColor(60, 65, 75))
            tf = m1.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(12)
            tf.margin_right = Pt(10)
            tf.margin_top = Pt(8)
            tf.margin_bottom = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.25
            r1 = p.add_run()
            r1.text = "動作："
            r1.font.name = FONT_NAME
            r1.font.size = Pt(12.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r2 = p.add_run()
            r2.text = " 刪除 .png, .css, .js, .woff 以及 GA、Facebook Pixel 等外部請求。"
            r2.font.name = FONT_NAME
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(203, 213, 225)

            # Purpose Box 1
            l, t, w, h = px(277, 448, 242, 138)
            b1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(b1, RGBColor(30, 32, 38), RGBColor(60, 65, 75))
            tf = b1.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(12)
            tf.margin_right = Pt(10)
            tf.margin_top = Pt(8)
            tf.margin_bottom = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.25
            r1 = p.add_run()
            r1.text = "目的："
            r1.font.name = FONT_NAME
            r1.font.size = Pt(12.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r2 = p.add_run()
            r2.text = " 避免浪費資源打爆外部服務\n"
            r2.font.name = FONT_NAME
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(203, 213, 225)
            r3 = p.add_run()
            r3.text = "(Don't load test Google!)\n"
            r3.font.name = FONT_NAME
            r3.font.size = Pt(12)
            r3.font.bold = True
            r3.font.color.rgb = RGBColor(52, 211, 153)
            r4 = p.add_run()
            r4.text = "並防止污染吞吐量報告。"
            r4.font.name = FONT_NAME
            r4.font.size = Pt(12)
            r4.font.color.rgb = RGBColor(203, 213, 225)

            # ---------------- CARD 2 (Dynamic Correlation) ----------------
            l, t, w, h = px(567, 225, 241, 68)
            h2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(h2, RGBColor(38, 41, 47), RGBColor(70, 75, 86))
            tf = h2.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = "② 動態 Token 關聯\n"
            r.font.name = FONT_NAME
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(241, 245, 249)
            r_sub = p.add_run()
            r_sub.text = "(Dynamic Correlation)"
            r_sub.font.name = FONT_NAME
            r_sub.font.size = Pt(11)
            r_sub.font.color.rgb = RGBColor(148, 163, 184)

            # Action Box 2
            l, t, w, h = px(567, 302, 241, 136)
            m2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(m2, RGBColor(30, 32, 38), RGBColor(60, 65, 75))
            tf = m2.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(12)
            tf.margin_right = Pt(10)
            tf.margin_top = Pt(8)
            tf.margin_bottom = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.25
            r1 = p.add_run()
            r1.text = "動作："
            r1.font.name = FONT_NAME
            r1.font.size = Pt(12.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r2 = p.add_run()
            r2.text = " 將寫死的 JWT 或 CSRF Token，改為從前置 API Response 動態提取。"
            r2.font.name = FONT_NAME
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(203, 213, 225)

            # Code Box 2
            l, t, w, h = px(567, 448, 241, 138)
            b2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(b2, RGBColor(18, 20, 24), RGBColor(50, 55, 65))
            tf = b2.text_frame
            tf.word_wrap = True
            tf.margin_left = Pt(16)
            tf.margin_top = Pt(10)

            p0 = tf.paragraphs[0]
            p0.alignment = PP_ALIGN.LEFT
            r = p0.add_run()
            r.text = "● ● ●"
            r.font.name = CODE_FONT
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(100, 116, 139)

            p1 = tf.add_paragraph()
            p1.alignment = PP_ALIGN.LEFT
            p1.line_spacing = 1.15
            r1 = p1.add_run()
            r1.text = "let "
            r1.font.name = CODE_FONT
            r1.font.size = Pt(12.5)
            r1.font.color.rgb = RGBColor(192, 132, 252)
            r2 = p1.add_run()
            r2.text = "token = "
            r2.font.name = CODE_FONT
            r2.font.size = Pt(12.5)
            r2.font.color.rgb = RGBColor(241, 245, 249)

            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.line_spacing = 1.15
            r3 = p2.add_run()
            r3.text = "res.json("
            r3.font.name = CODE_FONT
            r3.font.size = Pt(12.5)
            r3.font.color.rgb = RGBColor(253, 224, 71)

            p3 = tf.add_paragraph()
            p3.alignment = PP_ALIGN.LEFT
            p3.line_spacing = 1.15
            r4 = p3.add_run()
            r4.text = "  'accessToken');"
            r4.font.name = CODE_FONT
            r4.font.size = Pt(12.5)
            r4.font.color.rgb = RGBColor(52, 211, 153)

            # ---------------- CARD 3 (Groups & Checks) ----------------
            l, t, w, h = px(856, 225, 241, 68)
            h3 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(h3, RGBColor(38, 41, 47), RGBColor(70, 75, 86))
            tf = h3.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = "③ 補充結構化區塊\n"
            r.font.name = FONT_NAME
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(241, 245, 249)
            r_sub = p.add_run()
            r_sub.text = "(Groups & Checks)"
            r_sub.font.name = FONT_NAME
            r_sub.font.size = Pt(11)
            r_sub.font.color.rgb = RGBColor(148, 163, 184)

            # Action Box 3
            l, t, w, h = px(856, 302, 241, 136)
            m3 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(m3, RGBColor(30, 32, 38), RGBColor(60, 65, 75))
            tf = m3.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(12)
            tf.margin_right = Pt(10)
            tf.margin_top = Pt(8)
            tf.margin_bottom = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.25
            r1 = p.add_run()
            r1.text = "動作："
            r1.font.name = FONT_NAME
            r1.font.size = Pt(12.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r2 = p.add_run()
            r2.text = " 使用 "
            r2.font.name = FONT_NAME
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(203, 213, 225)
            r3 = p.add_run()
            r3.text = "group('Login') "
            r3.font.name = CODE_FONT
            r3.font.size = Pt(11)
            r3.font.bold = True
            r3.font.color.rgb = RGBColor(52, 211, 153)
            r4 = p.add_run()
            r4.text = "將 API 邏輯分組，並加上 "
            r4.font.name = FONT_NAME
            r4.font.size = Pt(12)
            r4.font.color.rgb = RGBColor(203, 213, 225)
            r5 = p.add_run()
            r5.text = "check() "
            r5.font.name = CODE_FONT
            r5.font.size = Pt(11)
            r5.font.bold = True
            r5.font.color.rgb = RGBColor(52, 211, 153)
            r6 = p.add_run()
            r6.text = "驗證狀態碼。"
            r6.font.name = FONT_NAME
            r6.font.size = Pt(12)
            r6.font.color.rgb = RGBColor(203, 213, 225)

            # Purpose Box 3
            l, t, w, h = px(856, 448, 241, 138)
            b3 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
            style_box(b3, RGBColor(30, 32, 38), RGBColor(60, 65, 75))
            tf = b3.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(12)
            tf.margin_right = Pt(10)
            tf.margin_top = Pt(8)
            tf.margin_bottom = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.25
            r1 = p.add_run()
            r1.text = "目的："
            r1.font.name = FONT_NAME
            r1.font.size = Pt(12.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r2 = p.add_run()
            r2.text = " 讓壓測報告具備真實業務意義，確保高負載下不但有回應，且業務邏輯完全正確。"
            r2.font.name = FONT_NAME
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(203, 213, 225)
            
        elif i == 8:
            s.shapes.add_picture(s8_clean, 0, 0, width=SW, height=SH)
            
            l, t, w, h = px(950, 67, 382, 102)
            box8 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
            style_box(box8, RGBColor(34, 38, 41), None)

            tf8 = box8.text_frame
            tf8.word_wrap = True
            tf8.margin_left = Pt(14)
            tf8.margin_right = Pt(12)
            tf8.margin_top = Pt(8)
            tf8.margin_bottom = Pt(6)
            tf8.vertical_anchor = MSO_ANCHOR.MIDDLE

            p1 = tf8.paragraphs[0]
            p1.alignment = PP_ALIGN.LEFT
            p1.line_spacing = 1.25
            r_title = p1.add_run()
            r_title.text = "解決痛點："
            r_title.font.name = FONT_NAME
            r_title.font.size = Pt(13.5)
            r_title.font.bold = True
            r_title.font.color.rgb = RGBColor(52, 211, 153)

            r_body1 = p1.add_run()
            r_body1.text = "純用 Browser 跑千人併發"
            r_body1.font.name = FONT_NAME
            r_body1.font.size = Pt(13)
            r_body1.font.color.rgb = RGBColor(241, 245, 249)

            p2 = tf8.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.line_spacing = 1.25
            r_body2 = p2.add_run()
            r_body2.text = "會瞬間吃爆壓測機成本；純用 API 跑"
            r_body2.font.name = FONT_NAME
            r_body2.font.size = Pt(13)
            r_body2.font.color.rgb = RGBColor(241, 245, 249)

            p3 = tf8.add_paragraph()
            p3.alignment = PP_ALIGN.LEFT
            p3.line_spacing = 1.25
            r_body3 = p3.add_run()
            r_body3.text = "則無法得知前端 UI 卡頓。"
            r_body3.font.name = FONT_NAME
            r_body3.font.size = Pt(13)
            r_body3.font.color.rgb = RGBColor(241, 245, 249)
            
        else:
            img_path = f'k6/slides/assets/Ch4/slide_{i}.png'
            s.shapes.add_picture(img_path, 0, 0, width=SW, height=SH)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    s4_clean, s5_clean, s8_clean = clean_slide_backgrounds()
    out_pptx = "k6/slides/Ch4_Precision_k6_Hybrid_Testing.pptx"
    build_presentation(s4_clean, s5_clean, s8_clean, out_pptx)
