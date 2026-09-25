#!/usr/bin/env python3
"""
Insert a native slide 8 "自訂指標的 5 個規則" into the current Ch3 deck (the Google
Slides export with embedded fonts), right after slide 7 (四大自訂指標型態).
Matches the codelab subsection and k6/demos/ch3_custom_metric_rules.js.

Idempotent: does nothing if a slide with this title already exists.
"""
import copy
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from add_metric_reading_slides import (  # noqa: E402
    CODELAB_URL, MONO, add_box, add_card_text, add_text, blank_layout, light_header, move_slide,
    L_ACCENT, L_BODY, L_BORDER, L_MUTED, L_TITLE, T_AMBER, T_DIM, T_GREEN, T_TEXT,
)
from build_ch3_native_deck import T, code_box  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "k6", "slides", "Ch3_k6_Quality_Gates.pptx")
TITLE = "自訂指標的 5 個規則"
INSERT_AT = 7  # 0-based: becomes slide 8, after 四大自訂指標型態


def badge_from(slide):
    for shp in slide.shapes:
        if shp._element.find(".//" + qn("a:hlinkClick")) is not None:
            return copy.deepcopy(shp._element)
    raise RuntimeError("badge not found")


def build(path=DECK):
    prs = Presentation(path)
    if any(TITLE in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame):
        print("custom-metric rules slide already present; nothing to do")
        return
    badge = badge_from(prs.slides[INSERT_AT - 1])

    s = prs.slides.add_slide(blank_layout(prs))
    light_header(s, TITLE + "：宣告、單位、命名、Tags、去哪裡看",
                 "四種型態會用了，接下來是定義與使用時最容易踩的坑（完整腳本：k6/demos/ch3_custom_metric_rules.js）")
    K, G, Y, W, D = "C4B5FD", T_GREEN, T_AMBER, T_TEXT, T_DIM
    code = [
        [("import", K), (" { Counter, Rate, Trend } ", W), ("from", K), (" 'k6/metrics';", G)],
        [("", W)],
        [("// ① 在最上層宣告  ③ 名稱只用英數與底線", D)],
        [("const", K), (" orders = ", W), ("new", K), (" Counter(", W), ("'orders_completed'", G), (");", W)],
        [("const", K), (" success = ", W), ("new", K), (" Rate(", W), ("'checkout_success'", G), (");", W)],
        [("// ② 時間型 Trend 加 true", D)],
        [("const", K), (" dbTime = ", W), ("new", K), (" Trend(", W), ("'db_query_time'", G), (", ", W), ("true", K), (");", W)],
        [("", W)],
        [("export const", K), (" options = { thresholds: {", W)],
        [("  // ④ 用 tags 設門檻", D)],
        [("  'db_query_time{endpoint:checkout}'", G), (": [", W), ("'p(95)<200'", Y), ("],", W)],
        [("} };", W)],
        [("", W)],
        [("export default function", K), (" () {", W)],
        [("  orders.add(1);", W)],
        [("  success.add(", W), ("true", K), (");", W)],
        [("  dbTime.add(elapsedMs, { endpoint: ", W), ("'checkout'", G), (" });", W)],
        [("}", W)],
    ]
    code_box(s, 60, 132, 640, 566, code, size=13.5, anchor=MSO_ANCHOR.MIDDLE, pad=(22, 12, 14, 12))

    rules = [
        ("①", "在最上層宣告", "寫在 default function 裡，k6 直接報錯：metrics must be declared in the init context"),
        ("②", "時間型 Trend 加 true", "沒加的話摘要只顯示 123.4，看不出單位；加了才會顯示 123.4ms"),
        ("③", "名稱只用英數與底線", "以字母或底線開頭；用中文命名會報錯：Invalid metric name"),
        ("④", "用 tags 細分並設門檻", ".add(值, { endpoint: 'checkout' }) 搭配 {endpoint:checkout} 門檻；tag 值別放動態資料"),
        ("⑤", "知道去哪裡看", "終端機：摘要的 CUSTOM 區塊。Prometheus：加 k6_ 前綴，Rate 加 _rate、Counter 加 _total"),
    ]
    y = 132
    for num, head, body in rules:
        add_box(s, 722, y, 594, 104, fill="FFFFFF", line=L_BORDER, radius=0.08)
        circ = add_box(s, 740, y + 28, 48, 48, fill=L_ACCENT, shape=MSO_SHAPE.OVAL)
        add_card_text(circ, [[(num, T(15, "FFFFFF", True))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        add_text(s, 802, y + 6, 500, 92, [
            [(head, T(14.5, L_TITLE, True))],
            [(body, T(12, L_BODY))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
        y += 104 + 11

    rid = s.part.relate_to(CODELAB_URL, RT.HYPERLINK, is_external=True)
    for h in badge.iter(qn("a:hlinkClick")):
        h.set(qn("r:id"), rid)
    s.shapes._spTree.append(badge)
    move_slide(prs, s, INSERT_AT)
    prs.save(path)
    print(f"inserted slide {INSERT_AT + 1} into {path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else DECK)
