#!/usr/bin/env python3
"""
Rebuild k6/slides/Ch5_k6_Observability_and_Modular_Architecture.pptx
1. Slide 1 to 3: Original with banner
2. Slide 4: Original + Callout on left (author's xk6 blog article)
3. Slide 5 to 8: Original with banner
4. Slide 9: Native PowerPoint text boxes and cards for "企業導入實作 Checklist (Takeaway)"
5. Slide 10: Native PowerPoint cards for "推薦延伸閱讀 (Author's Deep-Dive Articles)"
6. Slide 11: Native PowerPoint cards for "隨堂練習指引 (Hands-on Practice)"
7. All 11 slides have the clickable Codelabs banner at the bottom left.
"""
import os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SW = Emu(16256000)
SH = Emu(9144000)
EMU_PER_PX_X = 16256000 / 1376
EMU_PER_PX_Y = 9144000 / 768

def px(x, y, w, h):
    return (
        int(x * EMU_PER_PX_X),
        int(y * EMU_PER_PX_Y),
        int(w * EMU_PER_PX_X),
        int(h * EMU_PER_PX_Y)
    )

FONT_HEADING = "Noto Sans CJK TC"
FONT_BODY = "Noto Sans CJK TC"
FONT_CODE = "Noto Sans Mono CJK TC"

def style_box(shape, fill_rgb, line_rgb=None, line_width=1.5):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()

def add_codelab_banner(slide):
    l, t, w, h = px(40, 735, 510, 24)
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    style_box(badge, RGBColor(10, 20, 32), RGBColor(0, 180, 216), 1)
    tf = badge.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(int(8 * EMU_PER_PX_X))
    tf.margin_right = Emu(int(8 * EMU_PER_PX_X))
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    
    r1 = p.add_run()
    r1.text = "▶ 課程實戰 Codelabs："
    r1.font.name = FONT_HEADING
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)
    
    r2 = p.add_run()
    r2.text = "tedmax100.github.io/o11y_lab_for_dummies/k6-performance-testing/"
    r2.font.name = FONT_CODE
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(102, 252, 241)
    
    url = "https://tedmax100.github.io/o11y_lab_for_dummies/k6-performance-testing/"
    badge.click_action.hyperlink.address = url

def build_slide9(prs, bg_path, blank):
    s9 = prs.slides.add_slide(blank)
    s9.shapes.add_picture(bg_path, 0, 0, width=SW, height=SH)
    
    # Title Box
    l_title, t_title, w_title, h_title = px(60, 48, 1100, 65)
    title_box = s9.shapes.add_textbox(l_title, t_title, w_title, h_title)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_t = tf_title.paragraphs[0]
    p_t.alignment = PP_ALIGN.LEFT
    
    r_t1 = p_t.add_run()
    r_t1.text = "企業導入實作 Checklist "
    r_t1.font.name = FONT_HEADING
    r_t1.font.size = Pt(22)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(255, 255, 255)
    
    r_t2 = p_t.add_run()
    r_t2.text = "(Takeaway)"
    r_t2.font.name = FONT_HEADING
    r_t2.font.size = Pt(17)
    r_t2.font.color.rgb = RGBColor(102, 252, 241)
    
    p_sub = tf_title.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    r_sub = p_sub.add_run()
    r_sub.text = "將 k6 測試代碼化、SLO 門禁自動化與全鏈路可觀測性深植於團隊日常研發流程"
    r_sub.font.name = FONT_BODY
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(154, 160, 166)
    
    cards = [
        {
            "y": 125, "h": 165,
            "bg": RGBColor(18, 32, 46),
            "border": RGBColor(0, 180, 216),
            "header": "▶ 本地開發與團隊協作",
            "header_color": RGBColor(102, 252, 241),
            "items": [
                ("腳本即時調優", "壓測腳本開發與本地微調時，養成啟用 K6_WEB_DASHBOARD=true 即時觀測延遲與錯誤率變化的習慣。"),
                ("跨部門溝通基準", "交付團隊評估效能時，統一匯出 HTML Report 獨立報表作為跨部門對齊的效能基準線 (Baseline)。")
            ]
        },
        {
            "y": 305, "h": 165,
            "bg": RGBColor(36, 28, 20),
            "border": RGBColor(251, 140, 0),
            "header": "▶ CI/CD 流水線整合",
            "header_color": RGBColor(255, 171, 0),
            "items": [
                ("無人值守乾淨退場", "於自動化 CI 流水線中，嚴格配置 K6_WEB_DASHBOARD_PORT=-1 確保容器產出報表後正常退出。"),
                ("客製化 Runner 映像檔", "透過 xk6 搭配 Docker 確定性編譯，發布團隊專屬之 Custom k6 Runner 容器映像檔於 CI 倉庫。")
            ]
        },
        {
            "y": 485, "h": 200,
            "bg": RGBColor(26, 22, 38),
            "border": RGBColor(187, 134, 252),
            "header": "▶ 進階擴充與全域遙測",
            "header_color": RGBColor(187, 134, 252),
            "items": [
                ("底層中間件協定擴充", "評估微服務架構需求，透過 xk6 掛載 xk6-kafka / xk6-redis / xk6-sql 等底層中間件原生連線。"),
                ("時序數據串流對齊", "啟用 Prometheus Remote Write，將壓測指標帶上 Git Commit Tag 即時推入 Prometheus / InfluxDB。"),
                ("雙十字準星秒級定位", "於 Grafana 建立共享十字準星儀表板，將 k6 延遲與 K8s Pod CPU 限流 (CFS Throttling) 精準上下對齊。")
            ]
        }
    ]
    
    for c in cards:
        l_c, t_c, w_c, h_c = px(60, c["y"], 1255, c["h"])
        shape = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_c, t_c, w_c, h_c)
        style_box(shape, c["bg"], c["border"], 1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Emu(int(26 * EMU_PER_PX_X))
        tf.margin_right = Emu(int(26 * EMU_PER_PX_X))
        tf.margin_top = Emu(int(15 * EMU_PER_PX_Y))
        tf.margin_bottom = Emu(int(14 * EMU_PER_PX_Y))
        
        # Header
        p_h = tf.paragraphs[0]
        p_h.alignment = PP_ALIGN.LEFT
        r_h = p_h.add_run()
        r_h.text = c["header"]
        r_h.font.name = FONT_HEADING
        r_h.font.size = Pt(15.5)
        r_h.font.bold = True
        r_h.font.color.rgb = c["header_color"]
        p_h.space_after = Pt(10)
        
        # Items
        for term, desc in c["items"]:
            p_i = tf.add_paragraph()
            p_i.alignment = PP_ALIGN.LEFT
            r_icon = p_i.add_run()
            r_icon.text = "●  "
            r_icon.font.name = FONT_HEADING
            r_icon.font.size = Pt(10)
            r_icon.font.bold = True
            r_icon.font.color.rgb = c["header_color"]
            
            r_term = p_i.add_run()
            r_term.text = f"{term}： "
            r_term.font.name = FONT_HEADING
            r_term.font.size = Pt(12.5)
            r_term.font.bold = True
            r_term.font.color.rgb = RGBColor(255, 255, 255)
            
            r_desc = p_i.add_run()
            r_desc.text = desc
            r_desc.font.name = FONT_BODY
            r_desc.font.size = Pt(11.5)
            r_desc.font.color.rgb = RGBColor(215, 222, 235)
            p_i.space_after = Pt(7)
            
    add_codelab_banner(s9)

def build_slide10(prs, bg_path, blank):
    s10 = prs.slides.add_slide(blank)
    s10.shapes.add_picture(bg_path, 0, 0, width=SW, height=SH)
    
    # Title
    l_title, t_title, w_title, h_title = px(60, 48, 1100, 65)
    title_box = s10.shapes.add_textbox(l_title, t_title, w_title, h_title)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_t = tf_title.paragraphs[0]
    p_t.alignment = PP_ALIGN.LEFT
    
    r_t1 = p_t.add_run()
    r_t1.text = "推薦延伸閱讀 "
    r_t1.font.name = FONT_HEADING
    r_t1.font.size = Pt(22)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(255, 255, 255)
    
    r_t2 = p_t.add_run()
    r_t2.text = "(Author's Deep-Dive Articles)"
    r_t2.font.name = FONT_HEADING
    r_t2.font.size = Pt(17)
    r_t2.font.color.rgb = RGBColor(102, 252, 241)
    
    p_sub = tf_title.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    r_sub = p_sub.add_run()
    r_sub.text = "講師親撰之 k6 系列技術專欄 — 深入外掛生態系開發、前端真實體驗混合壓測與微服務全鏈路可觀測性閉環"
    r_sub.font.name = FONT_BODY
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(154, 160, 166)
    
    cards_data = [
        {
            "x": 60, "w": 395,
            "badge": "xk6 模組擴充 (Go Extension)",
            "badge_color": RGBColor(187, 134, 252),
            "border_color": RGBColor(124, 77, 255),
            "bg_color": RGBColor(25, 22, 38),
            "title": "Grafana xk6: 手把手從開發到編譯插件",
            "subtitle": "核心：Go-to-JS Bridge 與動態密碼插件實戰",
            "url_text": "ganhua.wang/grafana-xk6",
            "url": "https://ganhua.wang/grafana-xk6",
            "bullets": [
                ("Go-to-JS Bridge 架構", "深度解析 RootModule、ModuleInstance 與 modules.Register 註冊機制，掌握 Go 結構體映射至 JS 類別的核心原理。"),
                ("Web3 OTP 實戰案例", "以動態身份驗證為例，親手實作 k6/x/otp 擴充套件，在百萬併發壓測中動態生成一次性金鑰密碼與簽名。"),
                ("Docker 確定性編譯", "使用 grafana/xk6 官方容器抹平各 OS 開發環境依賴，在零配置前提下一鍵編譯出團隊專屬客製化 k6 二進位檔。")
            ]
        },
        {
            "x": 490, "w": 395,
            "badge": "前端混壓測試 (k6-browser)",
            "badge_color": RGBColor(0, 230, 118),
            "border_color": RGBColor(0, 180, 216),
            "bg_color": RGBColor(16, 28, 34),
            "title": "Grafana k6 瀏覽器測試實戰",
            "subtitle": "核心：Playwright 相容與 Core Web Vitals",
            "url_text": "ganhua.wang/grafana-k6-browser",
            "url": "https://ganhua.wang/grafana-k6-browser",
            "bullets": [
                ("Playwright 相容生態", "全面採用 chromium.launch()、page.goto() 與語意化 Locator，前端工程師無需轉換思維即可上手。"),
                ("BrowserContext 隔離", "在單一瀏覽器處理程序中實現多用戶完全獨立的 Session、Cookie 與 LocalStorage，大幅降低記憶體開銷。"),
                ("電商全流程自動化", "以 OpenTelemetry Demo 購物車為例，完整模擬加購物車與結帳，採集真實 Web Vitals 並自動截圖保存錯誤現場。")
            ]
        },
        {
            "x": 920, "w": 395,
            "badge": "全鏈路閉環 (Observability)",
            "badge_color": RGBColor(255, 171, 0),
            "border_color": RGBColor(251, 140, 0),
            "bg_color": RGBColor(32, 26, 18),
            "title": "Getting Started with Grafana k6",
            "subtitle": "核心：OpenTelemetry 串流與 GitLab CI",
            "url_text": "ganhua.wang/getting-started-with-grafana-k6",
            "url": "https://ganhua.wang/getting-started-with-grafana-k6-hands-on-practice",
            "bullets": [
                ("語意化測試組織", "運用 k6/http、check 與 group 模組化封裝業務邏輯，產出結構清晰且具備業務意義的階層化壓測報表。"),
                ("OTel 串流全面整合", "直連 OpenTelemetry Collector，將客戶端壓測指標與後端分散式 Traces、Logs 進行跨維度時間軸對齊。"),
                ("GitLab CI 自動化門禁", "將 k6 壓測無縫嵌入 CI/CD 流水線，以 Exit Code 99 自動攔截效能退化代碼，守護生產主幹穩定。")
            ]
        }
    ]
    
    for c in cards_data:
        l_c, t_c, w_c, h_c = px(c["x"], 130, c["w"], 505)
        card_shape = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_c, t_c, w_c, h_c)
        style_box(card_shape, c["bg_color"], c["border_color"], 1.5)
        
        tf_c = card_shape.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.TOP
        tf_c.margin_left = Emu(int(16 * EMU_PER_PX_X))
        tf_c.margin_right = Emu(int(16 * EMU_PER_PX_X))
        tf_c.margin_top = Emu(int(14 * EMU_PER_PX_Y))
        tf_c.margin_bottom = Emu(int(12 * EMU_PER_PX_Y))
        
        # Badge
        p_badge = tf_c.paragraphs[0]
        p_badge.alignment = PP_ALIGN.LEFT
        r_b = p_badge.add_run()
        r_b.text = f"● {c['badge']}"
        r_b.font.name = FONT_HEADING
        r_b.font.size = Pt(10.5)
        r_b.font.bold = True
        r_b.font.color.rgb = c["badge_color"]
        p_badge.space_after = Pt(4)
        
        # Title
        p_title = tf_c.add_paragraph()
        p_title.alignment = PP_ALIGN.LEFT
        r_t = p_title.add_run()
        r_t.text = c["title"]
        r_t.font.name = FONT_HEADING
        r_t.font.size = Pt(12.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(255, 255, 255)
        p_title.space_after = Pt(2)
        
        # Subtitle
        p_sub = tf_c.add_paragraph()
        p_sub.alignment = PP_ALIGN.LEFT
        r_s = p_sub.add_run()
        r_s.text = c["subtitle"]
        r_s.font.name = FONT_BODY
        r_s.font.size = Pt(9)
        r_s.font.color.rgb = RGBColor(160, 170, 185)
        p_sub.space_after = Pt(10)
        
        # Bullets
        for b_title, b_desc in c["bullets"]:
            p_b = tf_c.add_paragraph()
            p_b.alignment = PP_ALIGN.LEFT
            r_bt = p_b.add_run()
            r_bt.text = f"• {b_title}：\n  "
            r_bt.font.name = FONT_HEADING
            r_bt.font.size = Pt(9.5)
            r_bt.font.bold = True
            r_bt.font.color.rgb = RGBColor(245, 245, 250)
            
            r_bd = p_b.add_run()
            r_bd.text = b_desc
            r_bd.font.name = FONT_BODY
            r_bd.font.size = Pt(9)
            r_bd.font.color.rgb = RGBColor(185, 190, 200)
            p_b.space_after = Pt(8)
            
        # Button inside card
        l_btn, t_btn, w_btn, h_btn = px(c["x"] + 16, 585, c["w"] - 32, 36)
        btn_shape = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_btn, t_btn, w_btn, h_btn)
        style_box(btn_shape, RGBColor(32, 44, 62), c["border_color"], 1.5)
        tf_btn = btn_shape.text_frame
        tf_btn.word_wrap = False
        tf_btn.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_btn = tf_btn.paragraphs[0]
        p_btn.alignment = PP_ALIGN.CENTER
        r_btn = p_btn.add_run()
        r_btn.text = f"📖 閱讀文章 ➔ {c['url_text']}"
        r_btn.font.name = FONT_CODE
        r_btn.font.size = Pt(9.5)
        r_btn.font.bold = True
        r_btn.font.color.rgb = RGBColor(102, 252, 241)
        btn_shape.click_action.hyperlink.address = c["url"]
            
    # Bottom callout banner
    l_bot, t_bot, w_bot, h_bot = px(60, 648, 1255, 48)
    bot_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_bot, t_bot, w_bot, h_bot)
    style_box(bot_box, RGBColor(20, 28, 42), RGBColor(0, 180, 216), 1)
    tf_bot = bot_box.text_frame
    tf_bot.word_wrap = True
    tf_bot.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_bot.margin_left = Emu(int(14 * EMU_PER_PX_X))
    tf_bot.margin_right = Emu(int(14 * EMU_PER_PX_X))
    
    p_bot = tf_bot.paragraphs[0]
    p_bot.alignment = PP_ALIGN.LEFT
    r_bot1 = p_bot.add_run()
    r_bot1.text = "💡 學習價值： "
    r_bot1.font.name = FONT_HEADING
    r_bot1.font.size = Pt(10)
    r_bot1.font.bold = True
    r_bot1.font.color.rgb = RGBColor(255, 215, 0)
    
    r_bot2 = p_bot.add_run()
    r_bot2.text = "本課程奠定企業級效能工程體系，專欄文章則引領大家攻克真實業務中的前沿自訂挑戰，推薦課後務必搭配精讀！"
    r_bot2.font.name = FONT_BODY
    r_bot2.font.size = Pt(9.5)
    r_bot2.font.color.rgb = RGBColor(220, 225, 235)
    
    add_codelab_banner(s10)

def build_slide11(prs, bg_path, blank):
    s11 = prs.slides.add_slide(blank)
    s11.shapes.add_picture(bg_path, 0, 0, width=SW, height=SH)
    
    # Title Box
    l_title, t_title, w_title, h_title = px(60, 48, 1100, 65)
    title_box = s11.shapes.add_textbox(l_title, t_title, w_title, h_title)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_t = tf_title.paragraphs[0]
    p_t.alignment = PP_ALIGN.LEFT
    
    r_t1 = p_t.add_run()
    r_t1.text = "隨堂練習指引 "
    r_t1.font.name = FONT_HEADING
    r_t1.font.size = Pt(22)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(255, 255, 255)
    
    r_t2 = p_t.add_run()
    r_t2.text = "(Hands-on Practice)"
    r_t2.font.name = FONT_HEADING
    r_t2.font.size = Pt(17)
    r_t2.font.color.rgb = RGBColor(102, 252, 241)
    
    p_sub = tf_title.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    r_sub = p_sub.add_run()
    r_sub.text = "請依序完成三大實戰任務，體驗原生儀表板、CI/CD 離線報告匯出與 Docker 確定性編譯"
    r_sub.font.name = FONT_BODY
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(154, 160, 166)
    
    tasks = [
        {
            "x": 60, "w": 395,
            "badge": "任務一：啟動原生 Dashboard",
            "badge_color": RGBColor(102, 252, 241),
            "border_color": RGBColor(0, 180, 216),
            "bg_color": RGBColor(16, 28, 38),
            "code": "K6_WEB_DASHBOARD=true \\\nk6 run script.js",
            "results": [
                ("瀏覽器即時訪問", "打開瀏覽器訪問 http://127.0.0.1:5665"),
                ("動態曲線觀察", "觀察測試期間即時動態折線圖與 VUs 變化"),
                ("細部耗時拆解", "檢視各 HTTP 階段細部耗時曲線與門檻進度")
            ]
        },
        {
            "x": 490, "w": 395,
            "badge": "任務二：產出靜態 HTML 報告",
            "badge_color": RGBColor(187, 134, 252),
            "border_color": RGBColor(124, 77, 255),
            "bg_color": RGBColor(26, 22, 38),
            "code": "K6_WEB_DASHBOARD=true \\\nK6_WEB_DASHBOARD_PORT=-1 \\\nK6_WEB_DASHBOARD_EXPORT=report.html \\\nk6 run script.js",
            "results": [
                ("無人值守退出", "Port=-1 確保測試完成後 k6 行程自動乾淨退出"),
                ("產出靜態資產", "當前目錄自動生成單一檔案 report.html 報表"),
                ("跨團隊共享", "雙擊檔案即可離線檢視完整互動儀表板與數據")
            ]
        },
        {
            "x": 920, "w": 395,
            "badge": "任務三：初探 xk6 確定性編譯",
            "badge_color": RGBColor(255, 171, 0),
            "border_color": RGBColor(251, 140, 0),
            "bg_color": RGBColor(34, 26, 18),
            "code": "docker run --rm -u $(id -u):$(id -g) \\\n  -v $(pwd):/xk6 grafana/xk6 build latest \\\n  --with github.com/grafana/xk6-sql",
            "results": [
                ("抹平環境差異", "利用 Docker 官方映像消除本地 Go 開發環境依賴"),
                ("產出自訂引擎", "成功編譯整合 xk6-sql 模組之客製化 k6 執行檔"),
                ("版本功能驗證", "執行 ./k6 version 確認已成功封裝目標協定模組")
            ]
        }
    ]
    
    for t in tasks:
        # Card background shape
        l_t, t_t, w_t, h_t = px(t["x"], 125, t["w"], 565)
        card_shape = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_t, t_t, w_t, h_t)
        style_box(card_shape, t["bg_color"], t["border_color"], 1.5)
        
        # Header Box (Dedicated Shape)
        l_hb, t_hb, w_hb, h_hb = px(t["x"] + 14, 138, t["w"] - 28, 55)
        header_box = s11.shapes.add_textbox(l_hb, t_hb, w_hb, h_hb)
        tf_h = header_box.text_frame
        tf_h.word_wrap = True
        p_b = tf_h.paragraphs[0]
        p_b.alignment = PP_ALIGN.LEFT
        r_b = p_b.add_run()
        r_b.text = t["badge"]
        r_b.font.name = FONT_HEADING
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = t["badge_color"]
        p_b.space_after = Pt(4)
        
        p_l1 = tf_h.add_paragraph()
        p_l1.alignment = PP_ALIGN.LEFT
        r_l1 = p_l1.add_run()
        r_l1.text = "▶ 執行指令 (Command)："
        r_l1.font.name = FONT_HEADING
        r_l1.font.size = Pt(10.5)
        r_l1.font.bold = True
        r_l1.font.color.rgb = RGBColor(240, 240, 245)
        
        # Code Box (Dedicated Shape)
        l_code, t_code, w_code, h_code = px(t["x"] + 14, 200, t["w"] - 28, 115)
        code_box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_code, t_code, w_code, h_code)
        style_box(code_box, RGBColor(10, 15, 22), RGBColor(50, 65, 80), 1)
        tf_code = code_box.text_frame
        tf_code.word_wrap = True
        tf_code.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_code.margin_left = Emu(int(12 * EMU_PER_PX_X))
        tf_code.margin_right = Emu(int(12 * EMU_PER_PX_X))
        p_c = tf_code.paragraphs[0]
        p_c.alignment = PP_ALIGN.LEFT
        r_c = p_c.add_run()
        r_c.text = t["code"]
        r_c.font.name = FONT_CODE
        r_c.font.size = Pt(9.5)
        r_c.font.bold = True
        r_c.font.color.rgb = RGBColor(0, 230, 118)
        
        # Verification Box (Dedicated Shape)
        l_vb, t_vb, w_vb, h_vb = px(t["x"] + 14, 325, t["w"] - 28, 230)
        veri_box = s11.shapes.add_textbox(l_vb, t_vb, w_vb, h_vb)
        tf_v = veri_box.text_frame
        tf_v.word_wrap = True
        p_v = tf_v.paragraphs[0]
        p_v.alignment = PP_ALIGN.LEFT
        r_v = p_v.add_run()
        r_v.text = "▶ 驗證結果 (Verification)："
        r_v.font.name = FONT_HEADING
        r_v.font.size = Pt(10.5)
        r_v.font.bold = True
        r_v.font.color.rgb = RGBColor(255, 255, 255)
        p_v.space_after = Pt(8)
        
        for r_title, r_desc in t["results"]:
            p_r = tf_v.add_paragraph()
            p_r.alignment = PP_ALIGN.LEFT
            r_icon = p_r.add_run()
            r_icon.text = "● "
            r_icon.font.name = FONT_HEADING
            r_icon.font.size = Pt(10)
            r_icon.font.bold = True
            r_icon.font.color.rgb = t["badge_color"]
            
            r_head = p_r.add_run()
            r_head.text = f"{r_title}： "
            r_head.font.name = FONT_HEADING
            r_head.font.size = Pt(10)
            r_head.font.bold = True
            r_head.font.color.rgb = RGBColor(255, 255, 255)
            
            r_tail = p_r.add_run()
            r_tail.text = r_desc
            r_tail.font.name = FONT_BODY
            r_tail.font.size = Pt(9.5)
            r_tail.font.color.rgb = RGBColor(190, 195, 205)
            p_r.space_after = Pt(6)
            
        # Bottom pill: Demo script reference
        l_pill, t_pill, w_pill, h_pill = px(t["x"] + 14, 630, t["w"] - 28, 30)
        pill = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_pill, t_pill, w_pill, h_pill)
        style_box(pill, RGBColor(20, 30, 44), t["border_color"], 1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = False
        tf_pill.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_pill = tf_pill.paragraphs[0]
        p_pill.alignment = PP_ALIGN.CENTER
        r_pill = p_pill.add_run()
        r_pill.text = "▶ 配套演示腳本收錄於 k6/demos/ 目錄"
        r_pill.font.name = FONT_HEADING
        r_pill.font.size = Pt(8.5)
        r_pill.font.bold = True
        r_pill.font.color.rgb = RGBColor(102, 252, 241)
            
    add_codelab_banner(s11)

def build_presentation():
    bg_path = 'k6/slides/assets/Ch5/clean_tech_bg.png'
    asset_dir = 'k6/slides/assets/Ch5'
    
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]
    
    # 1. Slide 1 to 3: Original with banner
    for i in range(1, 4):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(f'{asset_dir}/image{i}.png', 0, 0, width=SW, height=SH)
        add_codelab_banner(s)
        
    # 2. Slide 4: Original + Callout on left
    s4 = prs.slides.add_slide(blank)
    s4.shapes.add_picture(f'{asset_dir}/image4.png', 0, 0, width=SW, height=SH)
    
    l, t, w, h = px(50, 475, 280, 215)
    callout = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    style_box(callout, RGBColor(22, 26, 38), RGBColor(140, 90, 255), 1.5)
    
    tf4 = callout.text_frame
    tf4.word_wrap = True
    tf4.vertical_anchor = MSO_ANCHOR.TOP
    tf4.margin_left = Emu(int(14 * EMU_PER_PX_X))
    tf4.margin_right = Emu(int(14 * EMU_PER_PX_X))
    tf4.margin_top = Emu(int(12 * EMU_PER_PX_Y))
    tf4.margin_bottom = Emu(int(12 * EMU_PER_PX_Y))
    
    p0 = tf4.paragraphs[0]
    r_badge = p0.add_run()
    r_badge.text = "▶ 講師深度實戰專欄"
    r_badge.font.name = FONT_HEADING
    r_badge.font.size = Pt(10.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(187, 134, 252)
    p0.space_after = Pt(5)
    
    p1 = tf4.add_paragraph()
    r_title = p1.add_run()
    r_title.text = "Grafana xk6: 手把手從開發到編譯出 k6 插件"
    r_title.font.name = FONT_HEADING
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(6)
    
    p2 = tf4.add_paragraph()
    r_desc = p2.add_run()
    r_desc.text = "• Web3 OTP 動態金鑰簽名插件\n• RootModule / ModuleInstance 解密\n• 突破原生極限支援私有協定"
    r_desc.font.name = FONT_BODY
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(195, 200, 210)
    p2.space_after = Pt(8)
    
    p3 = tf4.add_paragraph()
    r_url = p3.add_run()
    r_url.text = "▶ 閱讀專欄：ganhua.wang/grafana-xk6"
    r_url.font.name = FONT_CODE
    r_url.font.size = Pt(9.5)
    r_url.font.bold = True
    r_url.font.color.rgb = RGBColor(102, 252, 241)
    callout.click_action.hyperlink.address = "https://ganhua.wang/grafana-xk6"
    
    add_codelab_banner(s4)
    
    # 3. Slide 5 to 8: Original with banner
    for i in range(5, 9):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(f'{asset_dir}/image{i}.png', 0, 0, width=SW, height=SH)
        add_codelab_banner(s)
        
    # 4. Slide 9: Native Shapes for 企業導入實作 Checklist (Takeaway)
    build_slide9(prs, bg_path, blank)
    
    # 5. Slide 10: Native Shapes for 推薦延伸閱讀 (Author's Deep-Dive Articles)
    build_slide10(prs, bg_path, blank)
    
    # 6. Slide 11: Native Shapes for 隨堂練習指引 (Hands-on Practice)
    build_slide11(prs, bg_path, blank)
    
    output_path = 'k6/slides/Ch5_k6_Observability_and_Modular_Architecture.pptx'
    prs.save(output_path)
    print(f"Successfully rebuilt all {len(prs.slides)} slides into {output_path}")

if __name__ == '__main__':
    build_presentation()
