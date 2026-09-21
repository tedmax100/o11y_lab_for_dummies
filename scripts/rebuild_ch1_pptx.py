#!/usr/bin/env python3
"""
Insert the `k6 x agent` slide into `k6/slides/Ch1_Modern_Performance_Testing_with_k6.pptx` as Slide 5,
and also generate a standalone presentation `k6/slides/k6_x_agent.pptx`.
"""
import os
import pptx
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SW = Emu(16256000)
SH = Emu(9144000)
EMU_PER_PX_X = 16256000 / 1376
EMU_PER_PX_Y = 9144000 / 768

def px(x, y, w, h):
    return (
        int(x * EMU_PER_PX_X),
        int(y * EMU_PER_PX_Y),
        int(w * EMU_PER_PX_X),
        int(h * EMU_PER_PX_Y)
    )

SLIDE_IMG = 'k6/slides/assets/slide_k6_x_agent.png'
TARGET_URL = 'https://tedmax100.github.io/o11y_lab_for_dummies/k6-performance-testing/'
DISPLAY_URL = 'tedmax100.github.io/o11y_lab_for_dummies/k6-performance-testing/'

def add_hotspot(slide, x, y, w, h, url):
    l, t, width, height = px(x, y, w, h)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, width, height)
    shape.fill.background()
    shape.line.fill.background()
    shape.click_action.hyperlink.address = url
    return shape

def add_codelab_badge(slide):
    # Check if badge already exists
    for s in list(slide.shapes):
        if s.name == "Codelabs_Link_Badge":
            sp = s._element
            sp.getparent().remove(sp)

    left = Inches(0.55)
    top = Inches(9.70)
    width = Inches(6.60)
    height = Inches(0.24)
    
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.name = "Codelabs_Link_Badge"
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(255, 255, 255)
    badge.line.color.rgb = RGBColor(203, 213, 225)
    badge.line.width = Pt(1)
    badge.click_action.hyperlink.address = TARGET_URL

    tf = badge.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT

    r1 = p.add_run()
    r1.text = '▶ 課程實戰 Codelabs：'
    r1.font.name = 'Noto Sans CJK TC'
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(51, 65, 85)

    r2 = p.add_run()
    r2.text = DISPLAY_URL
    r2.font.name = 'Noto Sans CJK TC'
    r2.font.size = Pt(8.5)
    r2.font.bold = False
    r2.font.color.rgb = RGBColor(2, 132, 199)

def populate_agent_slide(slide):
    slide.shapes.add_picture(SLIDE_IMG, 0, 0, width=SW, height=SH)
    # Clickable hotspots for GitHub and docs URLs
    add_hotspot(slide, 160, 704, 380, 18, "https://github.com/grafana/xk6-subcommand-agent")
    add_hotspot(slide, 160, 720, 560, 18, "https://grafana.com/docs/k6/latest/set-up/configure-ai-assistant/bootstrap-with-k6-x-agent/")
    add_codelab_badge(slide)

def build_standalone_pptx():
    prs = pptx.Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    slide_layout = prs.slide_layouts[6] # blank
    slide = prs.slides.add_slide(slide_layout)
    populate_agent_slide(slide)
    
    out_path = 'k6/slides/k6_x_agent.pptx'
    prs.save(out_path)
    print(f"Created standalone PPTX: {out_path}")

def update_ch1_pptx():
    ch1_path = 'k6/slides/Ch1_Modern_Performance_Testing_with_k6.pptx'
    prs = pptx.Presentation(ch1_path)
    
    # Check if already 11 slides
    if len(prs.slides) == 11:
        print("Ch1 already has 11 slides. Updating Slide 5 (index 4)...")
        slide = prs.slides[4]
        # Remove existing shapes on slide 5
        for s in list(slide.shapes):
            sp = s._element
            sp.getparent().remove(sp)
        populate_agent_slide(slide)
    else:
        print(f"Ch1 currently has {len(prs.slides)} slides. Adding Slide 5...")
        slide_layout = prs.slide_layouts[6] # blank
        slide = prs.slides.add_slide(slide_layout)
        populate_agent_slide(slide)
        
        # Move new slide to index 4 (5th slide, right after Slide 4)
        slide_id_list = prs.slides._sldIdLst
        slide_id_list.insert(4, slide_id_list[-1])
        print(f"Ch1 updated to {len(prs.slides)} slides.")

    prs.save(ch1_path)
    print(f"Saved {ch1_path}")

if __name__ == '__main__':
    build_standalone_pptx()
    update_ch1_pptx()
