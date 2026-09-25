#!/usr/bin/env python3
"""
Insert a native slide 7 "何時仍該用 Closed Model" into the Ch2 deck, right after
slide 6 (scenarios 執行器範例). Slides 5-6 only show what goes wrong
with the closed model; this one shows when the closed model is still the right model.
Styled like the native slide 4 (build_ch2_stages_slide.py).

Idempotent: does nothing if a slide with this title already exists.
"""
import copy
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from add_metric_reading_slides import CODELAB_URL, MONO, S, add_box, add_card_text, add_text, move_slide  # noqa: E402
from build_ch2_stages_slide import (  # noqa: E402
    AMBER, BG, CARD, CARD_BORDER, CODE_BG, CYAN, MAGENTA, MUTED, STR, TEXT, WHITE,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "k6", "slides", "Ch2_Scientific_k6_Traffic_Modeling.pptx")
TITLE = "Open Model 不是萬用解：何時仍該用 Closed Model"
INSERT_AT = 6  # 0-based: becomes slide 7, after the scenarios (ramping-vus vs ramping-arrival-rate) slide

CASES = [
    ("固定數量的 client，等回應才送下一筆",
     "batch worker、MQ consumer、固定 50 位客服、IoT 輪詢——變慢時本來就會少送",
     "constant-vus"),
    ("要驗的是同時在線數／連線數",
     "1 萬條 WebSocket、session 上限、connection pool——arrival rate 管不到連線數",
     "ramping-vus"),
    ("Browser 測試（Ch4）",
     "每個 VU 都是一個 Chromium；open model 一遇到變慢就加開瀏覽器，先垮的是壓測機",
     "constant-vus"),
    ("測試資料只能用固定次數",
     "每個 VU 綁一組帳號、每筆訂單資料只能用一次——要的是精確的次數，不是速率",
     "per-vu-iterations"),
    ("Smoke 與共用環境初探",
     "1～2 個 VU 確認腳本能跑；變慢時自動降速，不會把大家共用的 staging 打爆",
     "vus: 1"),
]


def badge_from(slide):
    for shp in slide.shapes:
        if shp._element.find(".//" + qn("a:hlinkClick")) is not None:
            return copy.deepcopy(shp._element)
    raise RuntimeError("badge not found")


def build(path=DECK):
    prs = Presentation(path)
    if any(TITLE in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame):
        print("closed-model cases slide already present; nothing to do")
        return
    badge = badge_from(prs.slides[INSERT_AT - 1])

    s = prs.slides.add_slide(prs.slides[INSERT_AT - 2].slide_layout)
    add_box(s, 0, 0, 1376, 768, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 60, 26, 1256, 52, [[(TITLE, S(28, WHITE, True))]], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 60, 78, 1256, 28, [[("協調性漏測只在「真實流量本來就是 open」時才是問題；系統本身是 closed 的，硬用 open model 反而模擬錯了",
                                     S(14, CYAN))]], anchor=MSO_ANCHOR.MIDDLE)

    # left: the one question that decides the model
    top, bottom = 122, 690
    add_box(s, 60, top, 420, bottom - top, fill=CARD, line=CARD_BORDER, radius=0.04)
    add_text(s, 84, top + 18, 372, 30, [[("選模型只問一題", S(13, MUTED, True))]])
    add_text(s, 84, top + 52, 372, 84, [[("後端變慢時，真實的 client 會不會跟著少送？", S(20, WHITE, True))]],
             anchor=MSO_ANCHOR.TOP)
    branches = [
        (top + 150, "不會：使用者照樣湧入", "Open Model", "constant / ramping-arrival-rate", "公開網站、公開 API、活動搶購", CYAN),
        (top + 318, "會：client 等回應才送下一筆", "Closed Model", "constant-vus / ramping-vus / iterations", "右邊這 5 種情境", MAGENTA),
    ]
    for y, answer, model, executors, examples, color in branches:
        add_box(s, 84, y, 372, 150, fill=CODE_BG, line=color, radius=0.06, line_w=1.5)
        add_text(s, 104, y + 10, 336, 132, [
            [(answer, S(13, color, True))],
            [("→ " + model, S(19, WHITE, True))],
            [(executors, S(11.5, STR, False, MONO))],
            [(examples, S(12, TEXT))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    add_text(s, 84, top + 482, 372, 76, [
        [("⚠ 用 Closed 時記得：", S(12.5, AMBER, True)),
         ("量到的延遲依然偏樂觀，別拿它驗公開服務的 SLO", S(12.5, TEXT))],
    ], anchor=MSO_ANCHOR.MIDDLE)

    # right: when closed is the right model
    x, w, gap = 500, 816, 10
    rh = (bottom - top - gap * (len(CASES) - 1)) / len(CASES)
    y = top
    for i, (head, body, executor) in enumerate(CASES, 1):
        add_box(s, x, y, w, rh, fill=CARD, line=CARD_BORDER, radius=0.08)
        circ = add_box(s, x + 18, y + (rh - 40) / 2, 40, 40, fill=MAGENTA, shape=MSO_SHAPE.OVAL)
        add_card_text(circ, [[(str(i), S(15, BG, True))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        add_text(s, x + 74, y + 6, 540, rh - 12, [
            [(head, S(15, WHITE, True))],
            [(body, S(12.5, TEXT))],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=3)
        chip = add_box(s, x + w - 190, y + (rh - 36) / 2, 172, 36, fill=CODE_BG, line=CARD_BORDER, radius=0.2)
        add_card_text(chip, [[(executor, S(12.5, STR, False, MONO))]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        y += rh + gap

    rid = s.part.relate_to(CODELAB_URL, RT.HYPERLINK, is_external=True)
    for h in badge.iter(qn("a:hlinkClick")):
        h.set(qn("r:id"), rid)
    s.shapes._spTree.append(badge)
    move_slide(prs, s, INSERT_AT)
    prs.save(path)
    print(f"inserted slide {INSERT_AT + 1} into {path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else DECK)
